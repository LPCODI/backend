"""Document parser adapters for uploaded presentation material."""

from __future__ import annotations

import subprocess
import tempfile
from collections.abc import Iterable
from hashlib import sha1
from http import HTTPStatus
from pathlib import Path
from typing import Any, Protocol

from app.core import ApiError
from app.domain import ErrorCode
from app.services.interfaces import AnalysisInput, DocumentParseResult, JsonObject

PPTX_PARSER_PROVIDER = "python-pptx"
PPTX_PARSER_VERSION = "0.1.0"
PPTX_UNSUPPORTED_SOURCE_MESSAGE = "PPTX 파일만 이 파서에서 처리할 수 있습니다."
PPTX_DEPENDENCY_UNAVAILABLE_MESSAGE = "PPTX 파싱 의존성을 사용할 수 없습니다."
PPTX_FILE_NOT_FOUND_MESSAGE = "파싱할 PPTX 파일을 찾을 수 없습니다."
PDF_PARSER_PROVIDER = "pymupdf"
PDF_PARSER_VERSION = "0.1.0"
PDF_UNSUPPORTED_SOURCE_MESSAGE = "PDF 파일만 이 파서에서 처리할 수 있습니다."
PDF_DEPENDENCY_UNAVAILABLE_MESSAGE = "PDF 파싱 의존성을 사용할 수 없습니다."
PDF_FILE_NOT_FOUND_MESSAGE = "파싱할 PDF 파일을 찾을 수 없습니다."
PDF_ENCRYPTED_FILE_MESSAGE = "암호화된 PDF 파일은 파싱할 수 없습니다."
DOCX_PARSER_PROVIDER = "python-docx"
DOCX_PARSER_VERSION = "0.1.0"
DOCX_UNSUPPORTED_SOURCE_MESSAGE = "DOCX 파일만 이 파서에서 처리할 수 있습니다."
DOCX_DEPENDENCY_UNAVAILABLE_MESSAGE = "DOCX 파싱 의존성을 사용할 수 없습니다."
DOCX_FILE_NOT_FOUND_MESSAGE = "파싱할 DOCX 파일을 찾을 수 없습니다."
LIBREOFFICE_CONVERTER_PROVIDER = "libreoffice-headless"
LIBREOFFICE_CONVERTER_VERSION = "0.1.0"
LIBREOFFICE_UNSUPPORTED_SOURCE_MESSAGE = "PPT 또는 DOC 파일만 LibreOffice 변환으로 처리할 수 있습니다."
LIBREOFFICE_DEPENDENCY_UNAVAILABLE_MESSAGE = "LibreOffice 변환기를 사용할 수 없습니다."
LIBREOFFICE_CONVERSION_FAILED_MESSAGE = "LibreOffice 문서 변환에 실패했습니다."
PPT_PARSER_PROVIDER = f"{LIBREOFFICE_CONVERTER_PROVIDER}+{PPTX_PARSER_PROVIDER}"
PPT_PARSER_VERSION = "0.1.0"
PPT_UNSUPPORTED_SOURCE_MESSAGE = "PPT 파일만 이 파서에서 처리할 수 있습니다."
PPT_FILE_NOT_FOUND_MESSAGE = "파싱할 PPT 파일을 찾을 수 없습니다."
DOC_PARSER_PROVIDER = f"{LIBREOFFICE_CONVERTER_PROVIDER}+{DOCX_PARSER_PROVIDER}"
DOC_PARSER_VERSION = "0.1.0"
DOC_UNSUPPORTED_SOURCE_MESSAGE = "DOC 파일만 이 파서에서 처리할 수 있습니다."
DOC_FILE_NOT_FOUND_MESSAGE = "파싱할 DOC 파일을 찾을 수 없습니다."


class LibreOfficeCommandRunner(Protocol):
    def __call__(
        self,
        command: list[str],
        *,
        capture_output: bool,
        check: bool,
        text: bool,
        timeout: int,
    ) -> subprocess.CompletedProcess[str]:
        """Run a LibreOffice command."""


def _clean_text(value: str | None) -> str:
    if not value:
        return ""
    lines = (" ".join(line.split()) for line in value.splitlines())
    return "\n".join(line for line in lines if line)


def _split_title_and_body(raw_text: str) -> tuple[str | None, str]:
    lines = [line for line in raw_text.splitlines() if line.strip()]
    if not lines:
        return None, ""
    return lines[0], "\n".join(lines[1:])


def _shape_name(shape: Any) -> str | None:
    name = getattr(shape, "name", None)
    return str(name) if name else None


def _shape_position(shape: Any) -> JsonObject:
    return {
        "left": int(getattr(shape, "left", 0) or 0),
        "top": int(getattr(shape, "top", 0) or 0),
        "width": int(getattr(shape, "width", 0) or 0),
        "height": int(getattr(shape, "height", 0) or 0),
    }


def _iter_text_shapes(slide: Any, *, title_shape: Any | None) -> Iterable[str]:
    for shape in slide.shapes:
        if title_shape is not None and shape == title_shape:
            continue
        if getattr(shape, "has_text_frame", False):
            text = _clean_text(shape.text)
            if text:
                yield text


def _extract_tables(slide: Any) -> list[JsonObject]:
    tables: list[JsonObject] = []
    for shape_index, shape in enumerate(slide.shapes, start=1):
        if not getattr(shape, "has_table", False):
            continue
        rows = [
            [_clean_text(cell.text) for cell in row.cells]
            for row in shape.table.rows
        ]
        tables.append(
            {
                "shapeIndex": shape_index,
                "name": _shape_name(shape),
                "rowCount": len(rows),
                "columnCount": max((len(row) for row in rows), default=0),
                "rows": rows,
                "position": _shape_position(shape),
            }
        )
    return tables


def _extract_images(slide: Any) -> list[JsonObject]:
    images: list[JsonObject] = []
    for shape_index, shape in enumerate(slide.shapes, start=1):
        image = getattr(shape, "image", None)
        if image is None:
            continue
        images.append(
            {
                "shapeIndex": shape_index,
                "name": _shape_name(shape),
                "contentType": image.content_type,
                "extension": image.ext,
                "sha1": image.sha1,
                "sizeBytes": len(image.blob),
                "position": _shape_position(shape),
            }
        )
    return images


def _extract_chart_categories(chart: Any) -> list[str]:
    categories: list[str] = []
    for plot in getattr(chart, "plots", ()):
        for category in getattr(plot, "categories", ()):
            category_text = _clean_text(str(category))
            if category_text:
                categories.append(category_text)
        if categories:
            break
    return categories


def _extract_chart_series(chart: Any) -> list[JsonObject]:
    series_values: list[JsonObject] = []
    for series in getattr(chart, "series", ()):
        values = []
        for value in getattr(series, "values", ()):
            values.append(float(value) if isinstance(value, int | float) else value)
        series_values.append(
            {
                "name": _clean_text(str(getattr(series, "name", ""))) or None,
                "values": values,
            }
        )
    return series_values


def _extract_graphs(slide: Any) -> list[JsonObject]:
    graphs: list[JsonObject] = []
    for shape_index, shape in enumerate(slide.shapes, start=1):
        if not getattr(shape, "has_chart", False):
            continue
        chart = shape.chart
        chart_title = None
        if getattr(chart, "has_title", False):
            chart_title = _clean_text(chart.chart_title.text_frame.text)
        graphs.append(
            {
                "shapeIndex": shape_index,
                "name": _shape_name(shape),
                "chartType": str(chart.chart_type),
                "title": chart_title,
                "categories": _extract_chart_categories(chart),
                "series": _extract_chart_series(chart),
                "position": _shape_position(shape),
            }
        )
    return graphs


def _extract_notes_text(slide: Any) -> str | None:
    try:
        if not slide.has_notes_slide:
            return None
        notes_frame = slide.notes_slide.notes_text_frame
    except (AttributeError, ValueError):
        return None
    notes_text = _clean_text(notes_frame.text)
    return notes_text or None


class PptxDocumentParserService:
    """Extract normalized slide content from PowerPoint PPTX files."""

    provider = PPTX_PARSER_PROVIDER
    version = PPTX_PARSER_VERSION

    async def parse(self, source: AnalysisInput) -> DocumentParseResult:
        if source.file_path is None:
            raise ApiError(
                status_code=HTTPStatus.BAD_REQUEST,
                code=ErrorCode.PRESENTATION_FILE_UNSUPPORTED_TYPE,
                message=PPTX_UNSUPPORTED_SOURCE_MESSAGE,
            )

        pptx_path = Path(source.file_path)
        if pptx_path.suffix.lower() != ".pptx":
            raise ApiError(
                status_code=HTTPStatus.BAD_REQUEST,
                code=ErrorCode.PRESENTATION_FILE_UNSUPPORTED_TYPE,
                message=PPTX_UNSUPPORTED_SOURCE_MESSAGE,
                details={"filePath": str(pptx_path)},
            )
        if not pptx_path.exists():
            raise ApiError(
                status_code=HTTPStatus.NOT_FOUND,
                code=ErrorCode.PRESENTATION_FILE_NOT_FOUND,
                message=PPTX_FILE_NOT_FOUND_MESSAGE,
                details={"filePath": str(pptx_path)},
            )

        try:
            from pptx import Presentation as PptxPresentation
        except ImportError as exc:
            raise ApiError(
                status_code=HTTPStatus.INTERNAL_SERVER_ERROR,
                code=ErrorCode.DEPENDENCY_UNAVAILABLE,
                message=PPTX_DEPENDENCY_UNAVAILABLE_MESSAGE,
                details={"dependency": "python-pptx"},
            ) from exc

        presentation = PptxPresentation(str(pptx_path))
        slides: list[JsonObject] = []
        warnings: list[str] = []
        for slide_number, pptx_slide in enumerate(presentation.slides, start=1):
            title_shape = pptx_slide.shapes.title
            title = _clean_text(title_shape.text) if title_shape is not None else ""
            body_blocks = list(_iter_text_shapes(pptx_slide, title_shape=title_shape))
            notes_text = _extract_notes_text(pptx_slide)
            slides.append(
                {
                    "slideNumber": slide_number,
                    "sortOrder": slide_number,
                    "title": title or None,
                    "body": "\n\n".join(body_blocks),
                    "rawText": "\n\n".join(part for part in (title, *body_blocks) if part),
                    "notes": notes_text,
                    "tables": _extract_tables(pptx_slide),
                    "graphs": _extract_graphs(pptx_slide),
                    "images": _extract_images(pptx_slide),
                }
            )

        if not slides:
            warnings.append("PPTX file contains no slides.")

        return DocumentParseResult(
            provider=self.provider,
            version=self.version,
            slides=tuple(slides),
            warnings=tuple(warnings),
        )


class LibreOfficeDocumentConverter:
    """Convert legacy Office files with LibreOffice headless mode."""

    provider = LIBREOFFICE_CONVERTER_PROVIDER
    version = LIBREOFFICE_CONVERTER_VERSION

    def __init__(
        self,
        *,
        executable: str = "libreoffice",
        timeout_seconds: int = 60,
        command_runner: LibreOfficeCommandRunner = subprocess.run,
    ) -> None:
        self.executable = executable
        self.timeout_seconds = timeout_seconds
        self.command_runner = command_runner

    def convert(
        self,
        source_path: Path,
        *,
        output_dir: Path,
        target_extension: str,
    ) -> Path:
        """Convert a PPT or DOC file and return the generated file path."""

        normalized_source_path = Path(source_path)
        source_extension = normalized_source_path.suffix.lower()
        normalized_target_extension = target_extension.lower().lstrip(".")
        if source_extension not in {".ppt", ".doc"} or normalized_target_extension not in {
            "pptx",
            "docx",
        }:
            raise ApiError(
                status_code=HTTPStatus.BAD_REQUEST,
                code=ErrorCode.PRESENTATION_FILE_UNSUPPORTED_TYPE,
                message=LIBREOFFICE_UNSUPPORTED_SOURCE_MESSAGE,
                details={
                    "filePath": str(normalized_source_path),
                    "targetExtension": normalized_target_extension,
                },
            )

        output_dir.mkdir(parents=True, exist_ok=True)
        command = [
            self.executable,
            "--headless",
            "--convert-to",
            normalized_target_extension,
            "--outdir",
            str(output_dir),
            str(normalized_source_path),
        ]
        try:
            completed_process = self.command_runner(
                command,
                capture_output=True,
                check=False,
                text=True,
                timeout=self.timeout_seconds,
            )
        except FileNotFoundError as exc:
            raise ApiError(
                status_code=HTTPStatus.INTERNAL_SERVER_ERROR,
                code=ErrorCode.DEPENDENCY_UNAVAILABLE,
                message=LIBREOFFICE_DEPENDENCY_UNAVAILABLE_MESSAGE,
                details={"dependency": self.executable},
            ) from exc
        except subprocess.TimeoutExpired as exc:
            raise ApiError(
                status_code=HTTPStatus.INTERNAL_SERVER_ERROR,
                code=ErrorCode.DEPENDENCY_UNAVAILABLE,
                message=LIBREOFFICE_CONVERSION_FAILED_MESSAGE,
                details={
                    "filePath": str(normalized_source_path),
                    "timeoutSeconds": self.timeout_seconds,
                },
            ) from exc

        converted_path = output_dir / f"{normalized_source_path.stem}.{normalized_target_extension}"
        if completed_process.returncode != 0 or not converted_path.exists():
            raise ApiError(
                status_code=HTTPStatus.INTERNAL_SERVER_ERROR,
                code=ErrorCode.DEPENDENCY_UNAVAILABLE,
                message=LIBREOFFICE_CONVERSION_FAILED_MESSAGE,
                details={
                    "filePath": str(normalized_source_path),
                    "targetExtension": normalized_target_extension,
                    "returnCode": completed_process.returncode,
                    "stdout": completed_process.stdout,
                    "stderr": completed_process.stderr,
                },
            )
        return converted_path


class PptDocumentParserService:
    """Convert PPT files to PPTX, then parse with the PPTX adapter."""

    provider = PPT_PARSER_PROVIDER
    version = PPT_PARSER_VERSION

    def __init__(
        self,
        *,
        converter: LibreOfficeDocumentConverter | None = None,
        pptx_parser: PptxDocumentParserService | None = None,
    ) -> None:
        self.converter = converter or LibreOfficeDocumentConverter()
        self.pptx_parser = pptx_parser or PptxDocumentParserService()

    async def parse(self, source: AnalysisInput) -> DocumentParseResult:
        ppt_path = _validate_legacy_source(
            source=source,
            expected_extension=".ppt",
            unsupported_message=PPT_UNSUPPORTED_SOURCE_MESSAGE,
            file_not_found_message=PPT_FILE_NOT_FOUND_MESSAGE,
        )
        with tempfile.TemporaryDirectory(prefix="ai-speech-ppt-") as temp_dir:
            converted_path = self.converter.convert(
                ppt_path,
                output_dir=Path(temp_dir),
                target_extension=".pptx",
            )
            converted_result = await self.pptx_parser.parse(
                AnalysisInput(
                    resource_id=source.resource_id,
                    file_path=converted_path,
                    text=source.text,
                    metadata={
                        **source.metadata,
                        "convertedFrom": str(ppt_path),
                        "converter": self.converter.provider,
                    },
                )
            )
        return DocumentParseResult(
            provider=self.provider,
            version=self.version,
            slides=converted_result.slides,
            warnings=(
                f"Converted legacy PPT with {self.converter.provider}.",
                *converted_result.warnings,
            ),
        )


def _extract_pdf_text(page: Any) -> str:
    return _clean_text(page.get_text("text"))


def _validate_legacy_source(
    *,
    source: AnalysisInput,
    expected_extension: str,
    unsupported_message: str,
    file_not_found_message: str,
) -> Path:
    if source.file_path is None:
        raise ApiError(
            status_code=HTTPStatus.BAD_REQUEST,
            code=ErrorCode.PRESENTATION_FILE_UNSUPPORTED_TYPE,
            message=unsupported_message,
        )

    source_path = Path(source.file_path)
    if source_path.suffix.lower() != expected_extension:
        raise ApiError(
            status_code=HTTPStatus.BAD_REQUEST,
            code=ErrorCode.PRESENTATION_FILE_UNSUPPORTED_TYPE,
            message=unsupported_message,
            details={"filePath": str(source_path)},
        )
    if not source_path.exists():
        raise ApiError(
            status_code=HTTPStatus.NOT_FOUND,
            code=ErrorCode.PRESENTATION_FILE_NOT_FOUND,
            message=file_not_found_message,
            details={"filePath": str(source_path)},
        )
    return source_path


def _extract_pdf_images(page: Any) -> list[JsonObject]:
    images: list[JsonObject] = []
    page_dict = page.get_text("dict")
    for block_index, block in enumerate(page_dict.get("blocks", ()), start=1):
        if block.get("type") != 1:
            continue
        bbox = block.get("bbox") or (0, 0, 0, 0)
        images.append(
            {
                "blockIndex": block_index,
                "width": block.get("width"),
                "height": block.get("height"),
                "extension": block.get("ext"),
                "sizeBytes": len(block.get("image") or b""),
                "position": {
                    "left": int(bbox[0]),
                    "top": int(bbox[1]),
                    "width": int(bbox[2] - bbox[0]),
                    "height": int(bbox[3] - bbox[1]),
                },
            }
        )
    return images


class PdfDocumentParserService:
    """Extract page-based slide content from PDF files."""

    provider = PDF_PARSER_PROVIDER
    version = PDF_PARSER_VERSION

    async def parse(self, source: AnalysisInput) -> DocumentParseResult:
        if source.file_path is None:
            raise ApiError(
                status_code=HTTPStatus.BAD_REQUEST,
                code=ErrorCode.PRESENTATION_FILE_UNSUPPORTED_TYPE,
                message=PDF_UNSUPPORTED_SOURCE_MESSAGE,
            )

        pdf_path = Path(source.file_path)
        if pdf_path.suffix.lower() != ".pdf":
            raise ApiError(
                status_code=HTTPStatus.BAD_REQUEST,
                code=ErrorCode.PRESENTATION_FILE_UNSUPPORTED_TYPE,
                message=PDF_UNSUPPORTED_SOURCE_MESSAGE,
                details={"filePath": str(pdf_path)},
            )
        if not pdf_path.exists():
            raise ApiError(
                status_code=HTTPStatus.NOT_FOUND,
                code=ErrorCode.PRESENTATION_FILE_NOT_FOUND,
                message=PDF_FILE_NOT_FOUND_MESSAGE,
                details={"filePath": str(pdf_path)},
            )

        try:
            import fitz
        except ImportError as exc:
            raise ApiError(
                status_code=HTTPStatus.INTERNAL_SERVER_ERROR,
                code=ErrorCode.DEPENDENCY_UNAVAILABLE,
                message=PDF_DEPENDENCY_UNAVAILABLE_MESSAGE,
                details={"dependency": "pymupdf"},
            ) from exc

        with fitz.open(str(pdf_path)) as document:
            if document.needs_pass:
                raise ApiError(
                    status_code=HTTPStatus.BAD_REQUEST,
                    code=ErrorCode.PRESENTATION_FILE_UNSUPPORTED_TYPE,
                    message=PDF_ENCRYPTED_FILE_MESSAGE,
                    details={"filePath": str(pdf_path)},
                )

            slides: list[JsonObject] = []
            warnings: list[str] = []
            for page_index, page in enumerate(document, start=1):
                raw_text = _extract_pdf_text(page)
                title, body = _split_title_and_body(raw_text)
                slides.append(
                    {
                        "slideNumber": page_index,
                        "sortOrder": page_index,
                        "title": title,
                        "body": body,
                        "rawText": raw_text,
                        "notes": None,
                        "tables": [],
                        "graphs": [],
                        "images": _extract_pdf_images(page),
                    }
                )

        if not slides:
            warnings.append("PDF file contains no pages.")

        return DocumentParseResult(
            provider=self.provider,
            version=self.version,
            slides=tuple(slides),
            warnings=tuple(warnings),
        )


def _is_docx_heading(paragraph: Any) -> bool:
    style_name = getattr(getattr(paragraph, "style", None), "name", "") or ""
    return style_name in {"Title", "Heading 1", "Heading 2"} or style_name.startswith("Heading")


def _new_docx_slide(slide_number: int, title: str | None = None) -> JsonObject:
    return {
        "slideNumber": slide_number,
        "sortOrder": slide_number,
        "title": title,
        "bodyBlocks": [],
        "rawBlocks": [],
        "notes": None,
        "tables": [],
        "graphs": [],
        "images": [],
    }


def _finalize_docx_slide(slide: JsonObject) -> JsonObject:
    body_blocks = list(slide.pop("bodyBlocks"))
    raw_blocks = list(slide.pop("rawBlocks"))
    raw_text = "\n\n".join(raw_blocks)
    title = slide["title"]
    body = "\n\n".join(body_blocks)
    if title is None and raw_text:
        title, body = _split_title_and_body(raw_text)
    return {
        **slide,
        "title": title,
        "body": body,
        "rawText": raw_text,
    }


def _docx_table_rows(table: Any) -> list[list[str]]:
    return [[_clean_text(cell.text) for cell in row.cells] for row in table.rows]


def _extract_docx_paragraph_images(paragraph: Any) -> list[JsonObject]:
    images: list[JsonObject] = []
    for image_index, blip in enumerate(paragraph._p.xpath(".//*[local-name()='blip']"), start=1):
        relationship_id = blip.get(
            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
        )
        if not relationship_id:
            continue
        image_part = paragraph.part.related_parts.get(relationship_id)
        if image_part is None:
            continue
        blob = image_part.blob
        content_type = image_part.content_type
        extension = content_type.rsplit("/", maxsplit=1)[-1] if "/" in content_type else None
        images.append(
            {
                "imageIndex": image_index,
                "relationshipId": relationship_id,
                "contentType": content_type,
                "extension": extension,
                "sha1": sha1(blob).hexdigest(),
                "sizeBytes": len(blob),
            }
        )
    return images


class DocxDocumentParserService:
    """Extract slide-like sections from Word DOCX files."""

    provider = DOCX_PARSER_PROVIDER
    version = DOCX_PARSER_VERSION

    async def parse(self, source: AnalysisInput) -> DocumentParseResult:
        if source.file_path is None:
            raise ApiError(
                status_code=HTTPStatus.BAD_REQUEST,
                code=ErrorCode.PRESENTATION_FILE_UNSUPPORTED_TYPE,
                message=DOCX_UNSUPPORTED_SOURCE_MESSAGE,
            )

        docx_path = Path(source.file_path)
        if docx_path.suffix.lower() != ".docx":
            raise ApiError(
                status_code=HTTPStatus.BAD_REQUEST,
                code=ErrorCode.PRESENTATION_FILE_UNSUPPORTED_TYPE,
                message=DOCX_UNSUPPORTED_SOURCE_MESSAGE,
                details={"filePath": str(docx_path)},
            )
        if not docx_path.exists():
            raise ApiError(
                status_code=HTTPStatus.NOT_FOUND,
                code=ErrorCode.PRESENTATION_FILE_NOT_FOUND,
                message=DOCX_FILE_NOT_FOUND_MESSAGE,
                details={"filePath": str(docx_path)},
            )

        try:
            from docx import Document
            from docx.table import Table
            from docx.text.paragraph import Paragraph
        except ImportError as exc:
            raise ApiError(
                status_code=HTTPStatus.INTERNAL_SERVER_ERROR,
                code=ErrorCode.DEPENDENCY_UNAVAILABLE,
                message=DOCX_DEPENDENCY_UNAVAILABLE_MESSAGE,
                details={"dependency": "python-docx"},
            ) from exc

        document = Document(str(docx_path))
        slides: list[JsonObject] = []
        current_slide = _new_docx_slide(slide_number=1)

        for child in document.element.body.iterchildren():
            if child.tag.endswith("}p"):
                paragraph = Paragraph(child, document)
                paragraph_text = _clean_text(paragraph.text)
                paragraph_images = _extract_docx_paragraph_images(paragraph)
                if paragraph_text and _is_docx_heading(paragraph):
                    if current_slide["rawBlocks"] or current_slide["tables"] or current_slide["images"]:
                        slides.append(_finalize_docx_slide(current_slide))
                        current_slide = _new_docx_slide(slide_number=len(slides) + 1)
                    current_slide["title"] = paragraph_text
                    current_slide["rawBlocks"].append(paragraph_text)
                elif paragraph_text:
                    current_slide["bodyBlocks"].append(paragraph_text)
                    current_slide["rawBlocks"].append(paragraph_text)
                current_slide["images"].extend(paragraph_images)
                continue

            if child.tag.endswith("}tbl"):
                table = Table(child, document)
                rows = _docx_table_rows(table)
                current_slide["tables"].append(
                    {
                        "tableIndex": len(current_slide["tables"]) + 1,
                        "rowCount": len(rows),
                        "columnCount": max((len(row) for row in rows), default=0),
                        "rows": rows,
                    }
                )
                flattened_rows = [" | ".join(cell for cell in row if cell) for row in rows]
                table_text = "\n".join(row for row in flattened_rows if row)
                if table_text:
                    current_slide["rawBlocks"].append(table_text)

        if current_slide["rawBlocks"] or current_slide["tables"] or current_slide["images"]:
            slides.append(_finalize_docx_slide(current_slide))

        warnings: list[str] = []
        if not slides:
            warnings.append("DOCX file contains no parseable content.")

        return DocumentParseResult(
            provider=self.provider,
            version=self.version,
            slides=tuple(slides),
            warnings=tuple(warnings),
        )


class DocDocumentParserService:
    """Convert DOC files to DOCX, then parse with the DOCX adapter."""

    provider = DOC_PARSER_PROVIDER
    version = DOC_PARSER_VERSION

    def __init__(
        self,
        *,
        converter: LibreOfficeDocumentConverter | None = None,
        docx_parser: DocxDocumentParserService | None = None,
    ) -> None:
        self.converter = converter or LibreOfficeDocumentConverter()
        self.docx_parser = docx_parser or DocxDocumentParserService()

    async def parse(self, source: AnalysisInput) -> DocumentParseResult:
        doc_path = _validate_legacy_source(
            source=source,
            expected_extension=".doc",
            unsupported_message=DOC_UNSUPPORTED_SOURCE_MESSAGE,
            file_not_found_message=DOC_FILE_NOT_FOUND_MESSAGE,
        )
        with tempfile.TemporaryDirectory(prefix="ai-speech-doc-") as temp_dir:
            converted_path = self.converter.convert(
                doc_path,
                output_dir=Path(temp_dir),
                target_extension=".docx",
            )
            converted_result = await self.docx_parser.parse(
                AnalysisInput(
                    resource_id=source.resource_id,
                    file_path=converted_path,
                    text=source.text,
                    metadata={
                        **source.metadata,
                        "convertedFrom": str(doc_path),
                        "converter": self.converter.provider,
                    },
                )
            )
        return DocumentParseResult(
            provider=self.provider,
            version=self.version,
            slides=converted_result.slides,
            warnings=(
                f"Converted legacy DOC with {self.converter.provider}.",
                *converted_result.warnings,
            ),
        )


__all__ = [
    "DOC_FILE_NOT_FOUND_MESSAGE",
    "DOC_PARSER_PROVIDER",
    "DOC_PARSER_VERSION",
    "DOC_UNSUPPORTED_SOURCE_MESSAGE",
    "DOCX_DEPENDENCY_UNAVAILABLE_MESSAGE",
    "DOCX_FILE_NOT_FOUND_MESSAGE",
    "DOCX_PARSER_PROVIDER",
    "DOCX_PARSER_VERSION",
    "DOCX_UNSUPPORTED_SOURCE_MESSAGE",
    "LIBREOFFICE_CONVERSION_FAILED_MESSAGE",
    "LIBREOFFICE_CONVERTER_PROVIDER",
    "LIBREOFFICE_CONVERTER_VERSION",
    "LIBREOFFICE_DEPENDENCY_UNAVAILABLE_MESSAGE",
    "LIBREOFFICE_UNSUPPORTED_SOURCE_MESSAGE",
    "PDF_DEPENDENCY_UNAVAILABLE_MESSAGE",
    "PDF_ENCRYPTED_FILE_MESSAGE",
    "PDF_FILE_NOT_FOUND_MESSAGE",
    "PDF_PARSER_PROVIDER",
    "PDF_PARSER_VERSION",
    "PDF_UNSUPPORTED_SOURCE_MESSAGE",
    "PPT_FILE_NOT_FOUND_MESSAGE",
    "PPT_PARSER_PROVIDER",
    "PPT_PARSER_VERSION",
    "PPT_UNSUPPORTED_SOURCE_MESSAGE",
    "PPTX_DEPENDENCY_UNAVAILABLE_MESSAGE",
    "PPTX_FILE_NOT_FOUND_MESSAGE",
    "PPTX_PARSER_PROVIDER",
    "PPTX_PARSER_VERSION",
    "PPTX_UNSUPPORTED_SOURCE_MESSAGE",
    "DocDocumentParserService",
    "DocxDocumentParserService",
    "LibreOfficeDocumentConverter",
    "PdfDocumentParserService",
    "PptDocumentParserService",
    "PptxDocumentParserService",
]
