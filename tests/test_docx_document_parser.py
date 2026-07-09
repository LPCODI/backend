import asyncio
import base64
import tempfile
import unittest
from pathlib import Path

from app.core import ApiError
from app.domain import ErrorCode
from app.services import AnalysisInput, DocxDocumentParserService
from app.services.interfaces import DocumentParserService

try:
    from docx import Document
    from docx.shared import Inches
except ImportError:  # pragma: no cover - exercised only when document extra is absent.
    Document = None
    Inches = None


ONE_PIXEL_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/p9sAAAAASUVORK5CYII="
)


@unittest.skipIf(Document is None, "python-docx is not installed")
class DocxDocumentParserTest(unittest.TestCase):
    def test_docx_parser_extracts_heading_sections_text_tables_and_images(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            docx_path = Path(temp_dir) / "demo.docx"
            image_path = Path(temp_dir) / "pixel.png"
            image_path.write_bytes(ONE_PIXEL_PNG)
            document = Document()
            document.add_heading("Problem", level=1)
            document.add_paragraph("Users need structured practice.")
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Metric"
            table.cell(0, 1).text = "Value"
            table.cell(1, 0).text = "Latency"
            table.cell(1, 1).text = "120ms"
            document.add_picture(str(image_path), width=Inches(1))
            document.add_heading("Solution", level=1)
            document.add_paragraph("Provide deterministic feedback loops.")
            document.save(docx_path)

            result = asyncio.run(
                DocxDocumentParserService().parse(
                    AnalysisInput(resource_id="file-1", file_path=docx_path),
                )
            )

        self.assertIsInstance(DocxDocumentParserService(), DocumentParserService)
        self.assertEqual(result.provider, "python-docx")
        self.assertEqual(len(result.slides), 2)
        first_slide = result.slides[0]
        self.assertEqual(first_slide["slideNumber"], 1)
        self.assertEqual(first_slide["sortOrder"], 1)
        self.assertEqual(first_slide["title"], "Problem")
        self.assertIn("Users need structured practice.", first_slide["body"])
        self.assertIn("Latency | 120ms", first_slide["rawText"])
        self.assertEqual(first_slide["tables"][0]["rows"][1], ["Latency", "120ms"])
        self.assertEqual(first_slide["images"][0]["contentType"], "image/png")
        self.assertEqual(first_slide["graphs"], [])
        self.assertEqual(first_slide["notes"], None)
        self.assertEqual(result.slides[1]["title"], "Solution")

    def test_docx_parser_uses_first_paragraph_as_title_without_heading(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            docx_path = Path(temp_dir) / "plain.docx"
            document = Document()
            document.add_paragraph("Plain Title")
            document.add_paragraph("Plain body")
            document.save(docx_path)

            result = asyncio.run(
                DocxDocumentParserService().parse(
                    AnalysisInput(resource_id="file-1", file_path=docx_path),
                )
            )

        self.assertEqual(len(result.slides), 1)
        self.assertEqual(result.slides[0]["title"], "Plain Title")
        self.assertEqual(result.slides[0]["body"], "Plain body")

    def test_docx_parser_rejects_non_docx_source(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "demo.pdf"
            source_path.write_bytes(b"%PDF-1.4")

            with self.assertRaises(ApiError) as error_context:
                asyncio.run(
                    DocxDocumentParserService().parse(
                        AnalysisInput(resource_id="file-1", file_path=source_path),
                    )
                )

        self.assertEqual(error_context.exception.code, ErrorCode.PRESENTATION_FILE_UNSUPPORTED_TYPE)

    def test_docx_parser_rejects_missing_docx_source(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            missing_path = Path(temp_dir) / "missing.docx"

            with self.assertRaises(ApiError) as error_context:
                asyncio.run(
                    DocxDocumentParserService().parse(
                        AnalysisInput(resource_id="file-1", file_path=missing_path),
                    )
                )

        self.assertEqual(error_context.exception.code, ErrorCode.PRESENTATION_FILE_NOT_FOUND)


if __name__ == "__main__":
    unittest.main()
