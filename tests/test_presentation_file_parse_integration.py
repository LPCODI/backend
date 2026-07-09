import tempfile
import unittest
from collections.abc import Generator
from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.core import Settings
from app.db import Base, get_db
from app.domain import PresentationStatus
from app.main import create_app
from app.models import Presentation, PresentationFile, RefreshToken, Slide, User

try:
    from pptx import Presentation as PptxPresentation
except ImportError:  # pragma: no cover - exercised only when document extra is absent.
    PptxPresentation = None


@unittest.skipIf(PptxPresentation is None, "python-pptx is not installed")
class PresentationFileParseIntegrationTest(unittest.TestCase):
    def setUp(self) -> None:
        self.storage_dir = tempfile.TemporaryDirectory()
        self.engine = create_engine(
            "sqlite+pysqlite:///:memory:",
            connect_args={"check_same_thread": False},
            poolclass=StaticPool,
        )
        self.session_factory = sessionmaker(
            bind=self.engine,
            class_=Session,
            autocommit=False,
            autoflush=False,
            expire_on_commit=False,
        )
        Base.metadata.create_all(
            self.engine,
            tables=[
                User.__table__,
                RefreshToken.__table__,
                Presentation.__table__,
                PresentationFile.__table__,
                Slide.__table__,
            ],
        )
        self.settings = Settings(
            app_name="Presentation File Parse Integration Test",
            jwt_secret_key="presentation-file-parse-integration-test-secret",
            local_storage_path=Path(self.storage_dir.name),
            max_presentation_file_size_bytes=1024 * 1024,
            _env_file=None,
        )
        self.app = create_app(settings=self.settings)
        self.app.dependency_overrides[get_db] = self._override_db
        self.client = TestClient(self.app)

    def tearDown(self) -> None:
        self.app.dependency_overrides.clear()
        Base.metadata.drop_all(
            self.engine,
            tables=[
                Slide.__table__,
                PresentationFile.__table__,
                Presentation.__table__,
                RefreshToken.__table__,
                User.__table__,
            ],
        )
        self.engine.dispose()
        self.storage_dir.cleanup()

    def _override_db(self) -> Generator[Session, None, None]:
        with self.session_factory() as session:
            yield session

    def _signup_login_and_create_presentation(self) -> tuple[dict[str, str], int]:
        signup = self.client.post(
            "/api/v1/auth/signup",
            json={
                "email": "integration-presenter@example.com",
                "password": "correct horse battery staple",
                "name": "Integration Presenter",
            },
        )
        self.assertEqual(signup.status_code, 201)
        login = self.client.post(
            "/api/v1/auth/login",
            json={
                "email": "integration-presenter@example.com",
                "password": "correct horse battery staple",
            },
        )
        self.assertEqual(login.status_code, 200)
        headers = {"Authorization": f"Bearer {login.json()['data']['access_token']}"}
        create = self.client.post(
            "/api/v1/presentations",
            headers=headers,
            json={
                "title": "Upload Parse Integration",
                "total_duration_seconds": 600,
                "qa_duration_seconds": 120,
            },
        )
        self.assertEqual(create.status_code, 201)
        return headers, create.json()["data"]["presentation_id"]

    def _build_pptx_bytes(self) -> bytes:
        presentation = PptxPresentation()
        first_slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        first_slide.shapes.title.text = "Problem"
        first_slide.placeholders[1].text = "Existing practice tools stop at recording."
        first_slide.notes_slide.notes_text_frame.text = "Explain project motivation."

        second_slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        second_slide.shapes.title.text = "Solution"
        second_slide.placeholders[1].text = "AI feedback closes the preparation loop."

        output = BytesIO()
        presentation.save(output)
        return output.getvalue()

    def test_upload_parse_and_read_parse_result_persists_slides(self) -> None:
        headers, presentation_id = self._signup_login_and_create_presentation()
        pptx_bytes = self._build_pptx_bytes()

        upload = self.client.post(
            f"/api/v1/presentations/{presentation_id}/files",
            headers=headers,
            files={
                "file": (
                    "integration-demo.pptx",
                    pptx_bytes,
                    "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                )
            },
        )
        upload_payload = upload.json()
        self.assertEqual(upload.status_code, 201)
        file_id = upload_payload["data"]["file_id"]
        object_key = upload_payload["data"]["object_key"]
        self.assertTrue((Path(self.storage_dir.name) / object_key).exists())
        self.assertEqual(upload_payload["data"]["status"], "UPLOADED")

        parse = self.client.post(
            f"/api/v1/presentations/{presentation_id}/parse",
            headers=headers,
            json={"file_id": file_id},
        )
        parse_payload = parse.json()
        self.assertEqual(parse.status_code, 202)
        self.assertTrue(parse_payload["success"])
        self.assertEqual(parse_payload["data"]["file_id"], file_id)
        self.assertEqual(parse_payload["data"]["status"], "PARSED")
        self.assertEqual(parse_payload["data"]["slide_count"], 2)
        self.assertEqual(parse_payload["data"]["parser_provider"], "python-pptx")

        parse_result = self.client.get(
            f"/api/v1/presentations/{presentation_id}/parse-result",
            headers=headers,
        )
        result_payload = parse_result.json()
        self.assertEqual(parse_result.status_code, 200)
        self.assertEqual(result_payload["data"]["file"]["status"], "PARSED")
        self.assertEqual(result_payload["data"]["file"]["slide_count"], 2)
        self.assertEqual(result_payload["data"]["slide_count"], 2)
        self.assertEqual([slide["title"] for slide in result_payload["data"]["slides"]], ["Problem", "Solution"])
        self.assertEqual(result_payload["data"]["slides"][0]["notes_text"], "Explain project motivation.")
        self.assertIn("Existing practice tools", result_payload["data"]["slides"][0]["raw_text"])

        with self.session_factory() as session:
            presentation = session.get(Presentation, presentation_id)
            presentation_file = session.get(PresentationFile, file_id)
            slides = session.execute(select(Slide).order_by(Slide.sort_order)).scalars().all()

        self.assertEqual(presentation.status, PresentationStatus.PARSED.value)
        self.assertEqual(presentation_file.status, "PARSED")
        self.assertIsNotNone(presentation_file.parsed_at)
        self.assertEqual([slide.slide_number for slide in slides], [1, 2])
        self.assertEqual([slide.sort_order for slide in slides], [1, 2])
        self.assertFalse(any(slide.excluded for slide in slides))


if __name__ == "__main__":
    unittest.main()
