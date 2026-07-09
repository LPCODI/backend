import asyncio
import base64
import tempfile
import unittest
from pathlib import Path

from app.core import ApiError
from app.domain import ErrorCode
from app.services import AnalysisInput, PptxDocumentParserService

try:
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx import Presentation
    from pptx.util import Inches
except ImportError:  # pragma: no cover - exercised only when document extra is absent.
    CategoryChartData = None
    XL_CHART_TYPE = None
    Presentation = None
    Inches = None


ONE_PIXEL_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/p9sAAAAASUVORK5CYII="
)


@unittest.skipIf(Presentation is None, "python-pptx is not installed")
class PptxDocumentParserTest(unittest.TestCase):
    def test_pptx_parser_extracts_slide_text_table_notes_graphs_and_images(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            pptx_path = Path(temp_dir) / "demo.pptx"
            image_path = Path(temp_dir) / "pixel.png"
            image_path.write_bytes(ONE_PIXEL_PNG)
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = "Problem"
            slide.placeholders[1].text = "First point\nSecond point"
            table_shape = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(4), Inches(1))
            table_shape.table.cell(0, 0).text = "Metric"
            table_shape.table.cell(0, 1).text = "Value"
            table_shape.table.cell(1, 0).text = "Latency"
            table_shape.table.cell(1, 1).text = "120ms"
            chart_data = CategoryChartData()
            chart_data.categories = ["Before", "After"]
            chart_data.add_series("Confidence", (2.5, 4.0))
            chart = slide.shapes.add_chart(
                XL_CHART_TYPE.COLUMN_CLUSTERED,
                Inches(1),
                Inches(4.3),
                Inches(4),
                Inches(2),
                chart_data,
            ).chart
            chart.has_title = True
            chart.chart_title.text_frame.text = "Practice impact"
            slide.shapes.add_picture(str(image_path), Inches(5), Inches(1), width=Inches(1))
            slide.notes_slide.notes_text_frame.text = "Explain why latency matters."
            presentation.save(pptx_path)

            result = asyncio.run(
                PptxDocumentParserService().parse(
                    AnalysisInput(resource_id="file-1", file_path=pptx_path),
                )
            )

        self.assertEqual(result.provider, "python-pptx")
        self.assertEqual(len(result.slides), 1)
        parsed_slide = result.slides[0]
        self.assertEqual(parsed_slide["slideNumber"], 1)
        self.assertEqual(parsed_slide["sortOrder"], 1)
        self.assertEqual(parsed_slide["title"], "Problem")
        self.assertIn("First point", parsed_slide["body"])
        self.assertIn("Second point", parsed_slide["rawText"])
        self.assertEqual(parsed_slide["notes"], "Explain why latency matters.")
        self.assertEqual(parsed_slide["tables"][0]["rows"][1], ["Latency", "120ms"])
        self.assertEqual(parsed_slide["graphs"][0]["title"], "Practice impact")
        self.assertEqual(parsed_slide["graphs"][0]["categories"], ["Before", "After"])
        self.assertEqual(parsed_slide["graphs"][0]["series"][0]["name"], "Confidence")
        self.assertEqual(parsed_slide["graphs"][0]["series"][0]["values"], [2.5, 4.0])
        self.assertEqual(parsed_slide["images"][0]["contentType"], "image/png")

    def test_pptx_parser_rejects_non_pptx_source(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "demo.pdf"
            source_path.write_bytes(b"%PDF-1.4")

            with self.assertRaises(ApiError) as error_context:
                asyncio.run(
                    PptxDocumentParserService().parse(
                        AnalysisInput(resource_id="file-1", file_path=source_path),
                    )
                )

        self.assertEqual(error_context.exception.code, ErrorCode.PRESENTATION_FILE_UNSUPPORTED_TYPE)


if __name__ == "__main__":
    unittest.main()
